"""Spreadsheets (CSV, TSV, XLSX): rows keep their columns in search, exact calculations, charts."""
import datetime as dt
import io
import json
import pytest
from openpyxl import Workbook
from lapi_app import core, knowledge, spreadsheet, worker
from test_dataset_models import system, dataset
from test_agentic_chat import scripted, agentic, tool, thread_message
from test_agent_workspace import plan

CSV = 'Customer;Amount;Date;\nAcme;"1,234.50";2026-10-01;rush\nBeta;120;2026-11-15;\nAcme;(10);2026-11-20;\n'.encode()


def workbook():
    book = Workbook(); sheet = book.active; sheet.title = 'Sales'
    sheet.append(['Region', 'Month', 'Revenue'])
    for row in (['North', dt.date(2026, 1, 5), 100], [None, dt.date(2026, 2, 7), 150], ['South', dt.date(2026, 1, 9), 80], ['South', dt.date(2026, 2, 3), 40]):
        sheet.append(row)
    sheet.merge_cells('A2:A3')
    notes = book.create_sheet('Notes'); notes.append(['Topic', 'Owner']); notes.append(['Budget', 'Priya'])
    out = io.BytesIO(); book.save(out); return out.getvalue()


class Store(dict):
    def put(self, key, value): self[key] = value
    def get(self, key): return self[key]
    def put_json(self, key, value): self[key] = json.dumps(value).encode()
    def json(self, key): return json.loads(self[key])


def test_columns_types_header_and_merged_cells_are_read_carefully():
    sheet = spreadsheet.read('invoices.csv', CSV)[0]
    assert sheet['columns'] == ['Customer', 'Amount', 'Date', 'Column D'] and sheet['kinds'] == ['text', 'number', 'date', 'text']
    assert [n for n, _ in sheet['rows']] == [2, 3, 4]  # Spreadsheet row numbers, as the person sees them.
    sales, notes = spreadsheet.read('sales.xlsx', workbook())
    assert sales['rows'][1][1][0] == 'North' and 'merged cells were filled from their first cell' in sales['notes']
    assert notes['name'] == 'Notes' and notes['columns'] == ['Topic', 'Owner']
    with pytest.raises(spreadsheet.SpreadsheetError, match='old Excel'):
        spreadsheet.read('old.xlsx', b'\xd0\xcf\x11\xe0' + b'0' * 100)


def test_every_chunk_keeps_values_with_their_column_names_and_rows_are_never_cut():
    sheet = spreadsheet.read('invoices.csv', CSV)[0]
    table = [p for p in spreadsheet.pages([sheet]) if p.get('table_rows')][0]
    assert table['table_rows'][0][1] == 'Row 2 · Customer: Acme · Amount: 1,234.50 · Date: 2026-10-01 · Column D: rush'
    chunks = spreadsheet.pack(table['table_header'], table['table_rows'], lambda text: len(text) <= 160)
    assert len(chunks) == 3
    for text, first, last in chunks:
        assert text.startswith('Sheet "invoices" · columns: Customer, Amount, Date, Column D')
        assert all(line.startswith('Row ') and 'Customer: ' in line for line in text.splitlines()[1:])
    wide = spreadsheet.pack('Sheet "x" · columns: A, B', [(9, 'Row 9 · A: ' + 'a' * 60 + ' · B: ' + 'b' * 60)], lambda text: len(text) <= 100)
    assert [line.splitlines()[1].split(' · ')[0] for line, _, _ in wide] == ['Row 9 (part 1)', 'Row 9 (part 2)']
    summary = spreadsheet.summary(sheet)
    assert 'Amount: number, 3 filled, min -10, max 1234.5, total 1344.5' in summary


def test_a_merged_title_row_above_the_header_is_a_note_not_the_header():
    book = Workbook(); sheet = book.active; sheet.title = 'Budget'
    sheet.append(['FY2026 Budget']); sheet.merge_cells('A1:C1'); sheet.append(['Department', 'Q1', 'Q2']); sheet.append(['Engineering', 120, 135])
    out = io.BytesIO(); book.save(out)
    read = spreadsheet.read('budget.xlsx', out.getvalue())[0]
    assert read['columns'] == ['Department', 'Q1', 'Q2'] and 'title row: FY2026 Budget' in read['notes']
    assert read['rows'] == [(3, ['Engineering', 120, 135])]


def test_exact_calculations_group_filter_and_refuse_wrong_columns():
    store = Store(); spreadsheet.save(store, 'd', spreadsheet.read('sales.xlsx', workbook()))
    by_region = spreadsheet.query(store, 'd', {'group_by': ['Region'], 'metrics': [{'op': 'sum', 'column': 'Revenue'}], 'sort': {'by': 'sum of Revenue', 'desc': True}})
    assert by_region['rows'] == [['North', 250.0], ['South', 120.0]]
    by_month = spreadsheet.query(store, 'd', {'group_by': [{'column': 'Month', 'bucket': 'month'}], 'metrics': [{'op': 'sum', 'column': 'revenue'}, {'op': 'count'}], 'sort': {'by': 'Month (month)'}})
    assert by_month['rows'] == [['2026-01', 180.0, 2], ['2026-02', 190.0, 2]]
    listed = spreadsheet.query(store, 'd', {'filters': [{'column': 'Revenue', 'op': '>=', 'value': 100}], 'columns': ['Region', 'Revenue']})
    assert listed['rows'] == [['North', 100.0], ['North', 150.0]] and listed['row_numbers'] == [2, 3]
    # A derived figure (revenue = quantity × price, here Revenue × 2) is one metric, calculated row by row.
    doubled = spreadsheet.query(store, 'd', {'group_by': ['Region'], 'metrics': [{'op': 'sum', 'column': 'Revenue * 2'}]})
    assert doubled['columns'] == ['Region', 'sum of Revenue * 2'] and sorted(doubled['rows']) == [['North', 500.0], ['South', 240.0]]
    scaled = spreadsheet.query(store, 'd', {'columns': ['Region', 'Revenue / 100'], 'sort': {'by': 'Revenue / 100', 'desc': True}, 'limit': 1})
    assert scaled['rows'] == [['North', 1.5]]
    # The top row itself, not just the top value.
    best = spreadsheet.query(store, 'd', {'metrics': [{'op': 'max', 'column': 'Revenue * 2'}], 'columns': ['Region'], 'sort': {'by': 'Revenue * 2', 'desc': True}, 'limit': 1})
    assert best['columns'] == ['Region', 'Revenue * 2'] and best['rows'] == [['North', 300.0]] and best['row_numbers'] == [3]
    assert spreadsheet.query(store, 'd', {'sheet': 'Notes', 'filters': [{'column': 'owner', 'op': '=', 'value': 'PRIYA'}]})['matched_rows'] == 1
    for bad, message in (({'metrics': [{'op': 'sum', 'column': 'Region'}]}, 'not a number column'),
                         ({'filters': [{'column': 'Nope', 'op': '=', 'value': 1}]}, 'no column "Nope"'),
                         ({'sheet': 'Missing'}, 'No sheet named'),
                         ({'group_by': [{'column': 'Region', 'bucket': 'month'}]}, 'not a date column')):
        with pytest.raises(spreadsheet.SpreadsheetError, match=message):
            spreadsheet.query(store, 'd', bad)


def upload_sheet(s, d, name, raw, **data):
    return s.api.post('/api/documents/upload', files={'files': (name, raw, 'application/octet-stream')},
                      data={'dataset_id': d['id'], 'allow_external': 'true', **data})


def test_spreadsheets_upload_index_with_row_chunks_and_keep_the_exact_table(system):
    s = system; d = dataset(s)
    for name, raw in (('invoices.csv', CSV), ('sales.xlsx', workbook())):
        r = upload_sheet(s, d, name, raw); assert r.status_code == 200, r.text
        doc = r.json()['documents'][0]
        assert worker.run_once()
        assert s.api.get('/api/documents/' + doc['id']).json()['status'] == 'ready'
        chunks = core.store.json('documents/' + doc['id'] + '/chunks.json')['chunks']
        rows = [c for c in chunks if c.get('table')]
        assert rows and all(c['section'].startswith('Sheet: ') and c['row_start'] <= c['row_end'] for c in rows)
        assert any('summary' in c.get('section', '') for c in chunks)
        assert spreadsheet.tables(core.store, doc['id'])
    sales = [c for c in core.store.json('documents/' + doc['id'] + '/chunks.json')['chunks'] if c.get('sheet') == 'Sales' and c.get('table')]
    assert any('Row 3 · Region: North · Month: 2026-02-07 · Revenue: 150' in c['text'] for c in sales)


def test_bad_spreadsheets_and_merging_them_are_refused(system):
    s = system; d = dataset(s)
    renamed = upload_sheet(s, d, 'report.xlsx', b'%PDF-1.7 fake')
    assert renamed.status_code == 400 and 'not a real Excel' in renamed.json()['detail']
    old = upload_sheet(s, d, 'old.xlsx', b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'0' * 64)
    assert old.status_code == 400 and '.xls' in old.json()['detail']
    empty = upload_sheet(s, d, 'empty.csv', b'\n\n')
    assert empty.status_code == 400 and 'no data' in empty.json()['detail']
    merged = s.api.post('/api/documents/upload', files=[('files', ('a.csv', CSV, 'text/csv')), ('files', ('b.txt', b'notes', 'text/plain'))],
                        data={'dataset_id': d['id'], 'allow_external': 'true', 'merge_files': 'true'})
    assert merged.status_code == 422 and 'not merged' in merged.json()['detail']


def test_agent_answers_from_the_table_and_draws_a_chart(system, monkeypatch):
    s = system; d = dataset(s)
    doc = upload_sheet(s, d, 'sales.xlsx', workbook()).json()['documents'][0]
    assert worker.run_once(); knowledge.sync(s.owner, d['id']); agentic(monkeypatch)
    seen = scripted(monkeypatch,
        tool('query_table', document='sales.xlsx', group_by=[{'column': 'Month', 'bucket': 'month'}], metrics=[{'op': 'sum', 'column': 'Revenue'}],
             sort_by='Month (month)', chart={'type': 'line', 'title': 'Revenue by month', 'x': 'Month (month)', 'y': ['sum of Revenue']}),
        tool('query_table', document='sales.xlsx', metrics=[{'op': 'sum', 'column': 'Region'}]),
        *[{'text': 'Revenue rose from January to February, as the chart shows.', 'calls': []}] * 3)
    result = plan(s, d, text='Chart revenue by month')
    outputs = [json.loads(m['content']) for m in seen[-1]['messages'] if m.get('role') == 'tool' and m.get('tool_name') == 'query_table']
    assert outputs[0]['data']['rows'] == [['2026-01', 180.0], ['2026-02', 190.0]] and outputs[0]['citations']
    assert 'not a number column' in json.dumps(outputs[-1])
    assert 'query_table' in seen[0]['tools']
    message = thread_message(s, result)
    chart = (message.get('charts') or [None])[0]
    assert chart and chart['type'] == 'line' and chart['rows'] == [['2026-01', 180.0], ['2026-02', 190.0]] and chart['x'] == 'Month (month)'


# --- Extraction from spreadsheet rows ---------------------------------------------------------------

def ledger(rows=1200):
    book = Workbook(); sheet = book.active; sheet.title = 'Transactions'
    sheet.append(['Transaction ID', 'Date', 'Type', 'Counterparty', 'Amount', 'Amount (USD)', 'Status', 'Due Date'])
    for i in range(rows):
        sheet.append([f'TXN-{100001 + i}', dt.date(2025, 1, 1) + dt.timedelta(days=i % 365), 'Income' if i % 2 else 'Expense',
                      f'Vendor {i % 7}', round(10 + i * 1.25, 2), round((10 + i * 1.25) * 1.1, 2), 'Paid', dt.date(2025, 3, 1)])
    out = io.BytesIO(); book.save(out); return out.getvalue()


ROW_FIELDS = {'type': 'object', 'additionalProperties': False, 'required': ['transaction_id', 'transaction_date', 'transaction_type', 'amount', 'status', 'due_date', 'tax'],
              'properties': {'transaction_id': {'type': ['string', 'null'], 'description': 'Transaction ID: Extract the identifier exactly as written.'},
                             'transaction_date': {'type': ['string', 'null'], 'format': 'date', 'description': 'Transaction date: YYYY-MM-DD.'},
                             'transaction_type': {'type': ['string', 'null'], 'description': 'Transaction type: Income or Expense.'},
                             'amount': {'type': ['number', 'null'], 'description': 'Amount: plain number.'},
                             'status': {'type': ['string', 'null'], 'description': 'Payment status: as written.'},
                             'due_date': {'type': ['string', 'null'], 'format': 'date', 'description': 'Due date: YYYY-MM-DD.'},
                             'tax': {'type': ['number', 'null'], 'description': 'Tax: plain number.'}}}


def test_fields_match_columns_by_key_or_label():
    from lapi_app.lapi_agent import table_extraction
    columns = ['Transaction ID', 'Date', 'Type', 'Counterparty', 'Amount', 'Amount (USD)', 'Status', 'Due Date']
    assert table_extraction.match(ROW_FIELDS['properties'], columns) == {
        'transaction_id': 'Transaction ID', 'transaction_date': 'Date', 'transaction_type': 'Type',
        'amount': 'Amount', 'status': 'Status', 'due_date': 'Due Date'}
    nested = dict(ROW_FIELDS, properties={'items': {'type': 'array', 'items': {'type': 'string'}}})
    assert table_extraction.flat_fields(nested) is None


def test_large_spreadsheet_is_extracted_row_by_row_without_the_model(system, monkeypatch):
    s = system; d = dataset(s)
    doc = upload_sheet(s, d, 'ledger.xlsx', ledger()).json()['documents'][0]
    assert worker.run_once()
    from lapi_app import providers
    monkeypatch.setattr(providers, 'extract_with_evidence', lambda *a, **k: pytest.fail('the model must not be called for spreadsheet rows'))
    t = s.api.post('/api/templates', json={'name': 'Rows', 'schema': ROW_FIELDS, 'instructions': '', 'dataset_id': d['id']}).json()
    e = s.api.post('/api/extractions', json={'document_ids': [doc['id']], 'template_id': t['id'], 'dataset_id': d['id'], 'allow_external': True}).json()
    assert worker.run_once()
    x = s.api.get('/api/extractions/' + e['id']).json()
    assert x['status'] == 'ready', x.get('error')
    records = x['result']['records']
    assert len(records) == 1200 and x['table_rows'] is True
    assert records[0]['fields'] == {'transaction_id': 'TXN-100001', 'transaction_date': '2025-01-01', 'transaction_type': 'Expense',
                                    'amount': 10.0, 'status': 'Paid', 'due_date': '2025-03-01', 'tax': None}
    assert records[1199]['fields']['amount'] == round(10 + 1199 * 1.25, 2)
    assert records[4]['source_filename'] == 'ledger.xlsx · Transactions row 6'
    assert 'No column matched tax' in x['evidence_notice'] and '1,200 rows' in x['evidence_notice']
    assert any(q['field'] == '/records/4/fields' and q['quote'] == 'Row 6' for q in x['evidence'])


def test_prompt_without_matching_columns_still_uses_the_model(system, monkeypatch):
    s = system; d = dataset(s)
    doc = upload_sheet(s, d, 'ledger.xlsx', ledger(20)).json()['documents'][0]
    assert worker.run_once()
    schema = {'type': 'object', 'additionalProperties': False, 'required': ['total_income', 'largest_vendor'],
              'properties': {'total_income': {'type': ['number', 'null']}, 'largest_vendor': {'type': ['string', 'null']}}}
    t = s.api.post('/api/templates', json={'name': 'Totals', 'schema': schema, 'instructions': '', 'dataset_id': d['id']}).json()
    e = s.api.post('/api/extractions', json={'document_ids': [doc['id']], 'template_id': t['id'], 'dataset_id': d['id'], 'allow_external': True}).json()
    assert worker.run_once()
    x = s.api.get('/api/extractions/' + e['id']).json()
    assert x['status'] == 'ready' and not x.get('table_rows') and 'records' not in x['result']
