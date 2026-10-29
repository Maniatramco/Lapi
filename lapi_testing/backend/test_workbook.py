"""Workbooks are read completely and connected: cells, formulas, comments, links, colours, images,
charts, text boxes and embedded files, with a coverage report of anything not read."""
import datetime as dt
import io
import re
import zipfile
import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from lapi_app import core, ocr, workbook, worker
from test_dataset_models import system, dataset
from test_spreadsheets import upload_sheet

RECEIPT_TEXT = 'RideNow Cabs\nBill BILL-2025-0002\nTotal 245.37 USD'
TEXT_BOX = 'Approved by finance on 2025-02-01'


def receipt_png():
    from PIL import Image
    image = Image.new('RGB', (120, 60), 'white')
    out = io.BytesIO(); image.save(out, format='PNG'); return out.getvalue()


def rich_workbook(extra_parts=True):
    book = Workbook()
    tx = book.active; tx.title = 'Transactions'
    tx.append(['Transaction ID', 'Date', 'Counterparty', 'Reference No', 'Amount', 'FX Rate', 'Amount (USD)', 'Status'])
    rows = [('TXN-1', 'Acme Retail Ltd', 'INV-2025-0001', 1200.0, 1.0, 'Paid'),
            ('TXN-2', 'RideNow Cabs', 'BILL-2025-0002', 245.37, 1.0, 'Overdue'),
            ('TXN-3', 'Acme Retail Ltd', 'INV-2025-0003', 980.5, 1.1, 'Pending'),
            ('TXN-4', 'City Properties', 'BILL-2025-0004', 8000.0, 1.0, 'Paid'),
            ('TXN-5', 'RideNow Cabs', 'BILL-2025-0005', 99.99, 1.0, 'Paid'),
            ('TXN-6', 'City Properties', 'BILL-2025-0006', 8000.0, 1.0, 'Pending')]
    for i, (tid, who, ref, amount, fx, status) in enumerate(rows, start=2):
        tx.append([tid, dt.date(2025, 1, i), who, ref, amount, fx, f'=ROUND(E{i}*F{i},2)', status])
    tx['H3'].comment = Comment('Disputed: driver charged twice', 'Priya')
    tx['D2'].hyperlink = 'https://example.com/invoices/INV-2025-0001'
    tx['H4'].fill = PatternFill('solid', fgColor='FFFF0000')
    tx.row_dimensions[6].hidden = True
    rule = DataValidation(type='list', formula1='"Paid,Pending,Overdue"'); tx.add_data_validation(rule); rule.add('H2:H200')
    image = SheetImage(io.BytesIO(receipt_png())); tx.add_image(image, 'I3')
    vendors = book.create_sheet('Vendors')
    vendors.append(['Name', 'Country', 'Email'])
    for name, country in (('Acme Retail Ltd', 'Germany'), ('RideNow Cabs', 'India'), ('City Properties', 'India'), ('Unused Vendor', 'France')):
        vendors.append([name, country, name.split()[0].lower() + '@example.com'])
    summary = book.create_sheet('Summary')
    summary['A1'] = 'Summary'
    summary['A3'] = 'Total USD'; summary['B3'] = '=SUM(Transactions!G2:G7)'
    summary['A4'] = 'Vendors'; summary['B4'] = '=COUNTA(Vendors!A2:A5)'
    chart = BarChart(); chart.title = 'Amounts by transaction'
    chart.add_data(Reference(tx, min_col=5, min_row=1, max_row=7), titles_from_data=True)
    chart.set_categories(Reference(tx, min_col=1, min_row=2, max_row=7))
    summary.add_chart(chart, 'D2')
    out = io.BytesIO(); book.save(out)
    return inject(out.getvalue()) if extra_parts else out.getvalue()


def inject(raw):
    """Add what openpyxl cannot write: a text box, an embedded file, and parts Lapi does not read."""
    source = zipfile.ZipFile(io.BytesIO(raw)); out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == 'xl/drawings/drawing1.xml':
                box = ('<xdr:twoCellAnchor xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:from><xdr:col>9</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
                       '<xdr:to><xdr:col>12</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>7</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
                       '<xdr:sp><xdr:nvSpPr><xdr:cNvPr id="90" name="Note box"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr><xdr:spPr/>'
                       f'<xdr:txBody><a:bodyPr/><a:p><a:r><a:t>{TEXT_BOX}</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>')
                data = re.sub(rb'(</(?:\w+:)?wsDr>)', lambda m: box.encode() + m.group(1), data)
            if info.filename == 'xl/worksheets/sheet1.xml':
                ole = ('<oleObjects><oleObject xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" progId="Package" shapeId="1025" r:id="rIdEmb1">'
                       '<objectPr><anchor><from><xdr:col xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">9</xdr:col>'
                       '<xdr:colOff xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">0</xdr:colOff>'
                       '<xdr:row xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">4</xdr:row>'
                       '<xdr:rowOff xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">0</xdr:rowOff></from></anchor></objectPr></oleObject></oleObjects>')
                data = data.replace(b'</worksheet>', ole.encode() + b'</worksheet>')
            if info.filename == 'xl/worksheets/_rels/sheet1.xml.rels':
                data = data.replace(b'</Relationships>', b'<Relationship Id="rIdEmb1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" Target="../embeddings/contract-note.txt"/></Relationships>')
            target.writestr(info, data)
        target.writestr('xl/embeddings/contract-note.txt', 'Contract note: City Properties lease renews in March 2026.')
        target.writestr('xl/connections.xml', '<connections/>')
        target.writestr('xl/mystery/thing.xml', '<thing/>')
    return out.getvalue()


def stub_image(raw, item):
    return {'text': RECEIPT_TEXT, 'description': 'A taxi receipt from RideNow Cabs.'}


def read_file(name, data):
    return data.decode()


@pytest.fixture
def book():
    return workbook.read('ledger.xlsx', rich_workbook(), read_image=stub_image, read_file=read_file)


def sheet(book, name):
    return next(s for s in book['sheets'] if s['name'] == name)


def row_text(book, name, number):
    from lapi_app import spreadsheet
    s = sheet(book, name)
    return spreadsheet.record(s, number, dict(s['rows'])[number])


def test_formulas_excel_never_calculated_are_calculated_and_explained(book):
    tx = sheet(book, 'Transactions')
    assert dict(tx['rows'])[4][6] == pytest.approx(1078.55)                 # =ROUND(E4*F4,2)
    assert book['coverage']['read']['formulas calculated by Lapi'] == 8
    formula = book['column_formulas']['Transactions']['Amount (USD)']
    assert formula['formula'] == '=ROUND(E{r}*F{r},2)'.replace('{r}', 'n') and formula['meaning'] == '=ROUND(Amount*FX Rate,2)'
    total = next(e for e in book['lineage'] if e['sheet'] == 'Summary' and e['cell'] == 'B3')
    assert total['label'] == 'Total USD' and float(total['value']) == pytest.approx(1200 + 245.37 + 1078.55 + 8000 + 99.99 + 8000)
    link = next(l for l in book['links'] if l['kind'] == 'formula' and l['from'] == 'Summary' and l['to'] == 'Transactions')
    assert link['strength'] == 'proven' and 'SUM' in link['functions']


def test_comments_links_colours_and_hidden_rows_stay_with_their_row(book):
    assert 'note on Status: Disputed: driver charged twice (Priya)' in row_text(book, 'Transactions', 3)
    assert 'link on Reference No: https://example.com/invoices/INV-2025-0001' in row_text(book, 'Transactions', 2)
    assert 'Status highlighted red' in row_text(book, 'Transactions', 4)
    assert 'hidden row' in row_text(book, 'Transactions', 6)
    status = next(c for c in book['profiles'][0]['columns'] if c['name'] == 'Status')
    assert status['values'] == ['Paid', 'Overdue', 'Pending'] and status['allowed'] == 'Paid,Pending,Overdue'


def test_image_is_read_and_connected_to_its_row_by_place_and_by_value(book):
    image = next(i for i in book['items'] if i['kind'] == 'image')
    assert image['sheet'] == 'Transactions' and image['row'] == 3 and image['text'] == RECEIPT_TEXT
    ways = {(l['sheet'], l.get('row'), l['strength']): l for l in image['links']}
    assert ('Transactions', 3, 'proven') in ways                              # placed in row 3, and shows BILL-2025-0002 of row 3
    assert ways[('Transactions', 3, 'proven')]['agrees'] == ['Amount', 'Amount (USD)']
    assert any(l.get('value') == 'ridenow cabs' and l['strength'] == 'likely' for l in image['links'])
    assert 'image Transactions!I3' in row_text(book, 'Transactions', 3) and 'BILL-2025-0002' in row_text(book, 'Transactions', 3)


def test_a_receipt_that_disagrees_with_its_row_is_flagged():
    wrong = lambda raw, item: {'text': 'RideNow Cabs BILL-2025-0002 Total 254.37', 'description': ''}
    book = workbook.read('ledger.xlsx', rich_workbook(), read_image=wrong, read_file=read_file)
    image = next(i for i in book['items'] if i['kind'] == 'image')
    link = next(l for l in image['links'] if l.get('row') == 3)
    assert link['differs']['item'] == [254.37] and book['coverage']['mismatches']


def test_chart_text_box_and_embedded_file_are_read_and_anchored(book):
    chart = next(i for i in book['items'] if i['kind'] == 'chart')
    assert chart['sheet'] == 'Summary' and chart['chart']['title'] == 'Amounts by transaction'
    assert any(l['kind'] == 'chart' and l['to'] == 'Transactions' for l in book['links'])
    box = next(i for i in book['items'] if i['kind'] == 'text')
    assert box['text'] == TEXT_BOX and (box['sheet'], box['row']) == ('Transactions', 6)
    embedded = next(i for i in book['items'] if i['kind'] == 'embedded')
    assert embedded['name'] == 'contract-note.txt' and 'lease renews' in embedded['text'] and embedded['row'] == 5


def test_sheets_are_linked_by_shared_keys(book):
    link = next(l for l in book['links'] if l['kind'] == 'key' and l['from'] == 'Transactions' and l['to'] == 'Vendors')
    assert (link['from_column'], link['to_column'], link['strength']) == ('Counterparty', 'Name', 'proven')


def test_nothing_is_dropped_silently(book):
    coverage = book['coverage']
    assert {'what': 'unrecognised part', 'where': 'xl/mystery/thing.xml', 'reason': 'Lapi does not know this kind of content yet'} in coverage['not_read']
    assert any(n['what'] == 'data connections' for n in coverage['not_read']) and coverage['complete'] is False
    read = coverage['read']
    assert (read['images'], read['charts'], read['text boxes'], read['embedded files'], read['comments and notes']) == (1, 1, 1, 1, 1)
    text = workbook.coverage_text(coverage)
    assert 'Not read:' in text and 'xl/mystery/thing.xml' in text


def test_unreadable_image_is_reported_not_skipped():
    failing = lambda raw, item: {'error': 'OCR is disabled. Enable LAPI_OCR_ENABLED or upload a searchable PDF.'}
    book = workbook.read('ledger.xlsx', rich_workbook(extra_parts=False), read_image=failing)
    assert any(n['what'].startswith('image') and 'OCR is disabled' in n['reason'] and n['where'] == 'Transactions!I3' for n in book['coverage']['not_read'])


def test_csv_lines_with_missing_cells_are_reported():
    book = workbook.read('a.csv', b'Name,Amount,Note\nAcme,10,ok\nBeta,20\n')
    assert any('fewer cells' in n for n in book['coverage']['notes'])


def test_password_protected_workbook_has_a_clear_message():
    from lapi_app import spreadsheet
    raw = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 64 + 'EncryptedPackage'.encode('utf-16-le')
    with pytest.raises(spreadsheet.SpreadsheetError, match='password protected'):
        spreadsheet.read('locked.xlsx', raw)


def test_indexing_keeps_every_part_searchable_and_saves_the_map(system, monkeypatch):
    monkeypatch.setattr(ocr, 'read_image', lambda raw, **k: {'text': RECEIPT_TEXT, 'text_origin': 'ocr', 'ocr': {'model': 'test-vision', 'provider': 'ollama', 'verified': False}})
    monkeypatch.setattr(ocr, 'describe', lambda raw, **k: 'A taxi receipt from RideNow Cabs.')
    s = system; d = dataset(s)
    doc = upload_sheet(s, d, 'ledger.xlsx', rich_workbook()).json()['documents'][0]
    assert worker.run_once()
    assert s.api.get('/api/documents/' + doc['id']).json()['status'] == 'ready'
    chunks = core.store.json('documents/' + doc['id'] + '/chunks.json')['chunks']
    text = '\n'.join(c['text'] for c in chunks)
    for expected in ('Disputed: driver charged twice', 'Amounts by transaction', TEXT_BOX, 'lease renews', 'Meaning profile of sheet "Vendors"',
                     'What Lapi read in this workbook', 'Total 245.37', 'Summary!B3 ("Total USD")'):
        assert expected in text, expected
    image = next(c for c in chunks if c.get('workbook_item', '').startswith('image'))
    assert image['ocr']['model'] == 'test-vision' and image['sheet'] == 'Transactions'
    saved = workbook.load(core.store, doc['id'])
    assert saved['coverage']['read']['images'] == 1 and any(l['kind'] == 'key' for l in saved['links'])


def test_excel_functions_pycel_lacks():
    from lapi_app import excel_functions as f
    keys = (('A',), ('B',), ('C',)); values = ((10,), (20,), (30,))
    assert f.xlookup('b', keys, values) == 20 and f.xlookup('z', keys, values, 'none') == 'none'
    assert f.xmatch(25, ((10,), (20,), (30,)), -1) == 2 and f.xmatch(25, ((10,), (20,), (30,)), 1) == 3
    assert f.counta(((1, None), ('', 'x'))) == 2 and f.countblank(((1, None), ('', 'x'))) == 2
    assert f.textjoin(', ', True, (('a',), (None,), ('b',))) == 'a, b'
    start, end = 45658, 45688          # 2025-01-01 (Wednesday) .. 2025-01-31
    assert f.networkdays(start, end) == 23 and f.workday(start, 2) == 45660
    assert f.datedif(45658, 46022, 'M') == 11 and f.datedif(45658, 46022, 'D') == 364
    assert f.switch(2, 1, 'one', 2, 'two', 'other') == 'two' and f.switch(9, 1, 'one', 'other') == 'other'
    assert f.median((3, 1, 2)) == 2 and f.rank(20, ((10,), (20,), (30,))) == 2 and f.proper('city properties ltd') == 'City Properties Ltd'
