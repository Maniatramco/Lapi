"""Spreadsheets (CSV, TSV, XLSX) as tables that keep every value with its column.

Three uses, one parsed form:
- Search: each data row becomes a self-contained record ("Row 5 · Customer: Acme · Amount: 450"), and
  rows are packed into chunks only at row boundaries, each chunk starting with the sheet and its
  columns (`pack`). A value is never separated from its column name or row number.
- Overview: one summary per sheet (columns, types, counts, ranges, examples) so questions about the
  whole sheet find it.
- Exact answers: the typed table is stored as Parquet next to the document (`save`), and `query`
  runs a fixed, validated set of operations on it (filter, group, count/sum/avg/min/max). The model
  chooses the operation; it never writes code or SQL, and numbers come from the data, not the model.
"""
import csv
import datetime as dt
import io
import re
from decimal import Decimal

EXTENSIONS = ('csv', 'tsv', 'xlsx')
MAX_ROWS = 50_000          # data rows across all sheets
MAX_COLUMNS = 200
MAX_CELL = 2_000           # characters kept per cell
MAX_SHEETS = 20
SUMMARY_EXAMPLES = 5


class SpreadsheetError(ValueError):
    """A spreadsheet Lapi cannot read, with a message the person can act on."""


def is_spreadsheet(name):
    return name.rsplit('.', 1)[-1].lower() in EXTENSIONS if '.' in name else False


# --- Reading ---------------------------------------------------------------------------------------

def _decode(raw):
    for encoding in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SpreadsheetError('The file is not readable text. Save it as UTF-8 CSV or as .xlsx.')


def _csv_rows(name, raw):
    text = _decode(raw)
    if '\x00' in text[:4096]:
        raise SpreadsheetError('This does not look like a CSV file. Save it as CSV (comma separated) or .xlsx.')
    if name.lower().endswith('.tsv'):
        delimiter = '\t'
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:20_000], delimiters=',;\t|').delimiter
        except csv.Error:
            delimiter = ','
    return [row for row in csv.reader(io.StringIO(text, newline=''), delimiter=delimiter)]


def _xlsx_sheets(raw):
    """(sheet name, rows, notes) for every sheet; merged ranges filled from their top-left cell."""
    if not raw.startswith(b'PK\x03\x04'):
        if raw.startswith(b'\xd0\xcf\x11\xe0') and 'EncryptedPackage'.encode('utf-16-le') in raw[:1_000_000]:
            raise SpreadsheetError('The workbook is password protected. Remove the password in Excel (File > Info > Protect Workbook) and upload it again.')
        if raw.startswith(b'\xd0\xcf\x11\xe0'):
            raise SpreadsheetError('This is an old Excel (.xls) file. Open it in Excel and save it as .xlsx or CSV.')
        raise SpreadsheetError('This is not an .xlsx workbook. Save it from Excel as .xlsx or CSV.')
    from openpyxl import load_workbook
    try:
        # data_only: the values Excel saved for formulas, not the formulas themselves.
        book = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
    except Exception as exc:
        raise SpreadsheetError('The workbook could not be opened. Save it again from Excel as .xlsx.') from exc
    sheets = []
    for sheet in book.worksheets[:MAX_SHEETS + 1]:
        notes = []
        if sheet.sheet_state != 'visible':
            notes.append('hidden sheet')
        filled = {}
        for merged in list(sheet.merged_cells.ranges):
            value = sheet.cell(merged.min_row, merged.min_col).value
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    filled[(row, col)] = value
        if filled:
            notes.append('merged cells were filled from their first cell')
        rows = []
        for r, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            rows.append([filled.get((r, c), value) for c, value in enumerate(row, start=1)])
        sheets.append((sheet.title, rows, notes))
    if len(book.worksheets) > MAX_SHEETS:
        raise SpreadsheetError(f'The workbook has more than {MAX_SHEETS} sheets. Split it or remove unused sheets.')
    return sheets


def _empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _trim(rows):
    """Drop empty rows at the end and empty columns on the right."""
    rows = [list(row) for row in rows]
    while rows and all(_empty(v) for v in rows[-1]):
        rows.pop()
    width = max((max((i + 1 for i, v in enumerate(row) if not _empty(v)), default=0) for row in rows), default=0)
    return [row[:width] + [None] * (width - len(row)) for row in rows], width


def _column_letter(index):
    letters = ''
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


def _headers(first, width):
    """Column names from the header row: blanks become "Column C", duplicates get "(2)"."""
    names, seen = [], {}
    for i in range(width):
        value = first[i] if i < len(first) else None
        name = ' '.join(str(value).split())[:120] if not _empty(value) else f'Column {_column_letter(i)}'
        count = seen.get(name.lower(), 0) + 1
        seen[name.lower()] = count
        names.append(name if count == 1 else f'{name} ({count})')
    return names


def _looks_like_header(row, below):
    """A header row is mostly text, and the rows under it are not all text in the same places."""
    cells = [v for v in row if not _empty(v)]
    if not cells or sum(isinstance(v, str) for v in cells) < max(1, len(cells) * 0.7):
        return False
    if not below:
        return True
    numeric_below = any(_number(v) is not None for r in below[:5] for v in r if not _empty(v))
    return numeric_below or all(isinstance(v, str) for v in cells)


NUMBER = re.compile(r'^[-+]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$')


def _number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(' ', '')
        for symbol in ('$', '€', '£', '₹', '¥'):
            text = text.replace(symbol, '')
        negative = text.startswith('(') and text.endswith(')')
        text = text.strip('()').strip()
        if NUMBER.match(text):
            number = float(text.replace(',', ''))
            return -number if negative else number
    return None


DATE = re.compile(r'^\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?$')


def _date(value):
    if isinstance(value, dt.datetime):
        return value
    if isinstance(value, dt.date):
        return dt.datetime(value.year, value.month, value.day)
    if isinstance(value, str) and DATE.match(value.strip()):
        try:
            return dt.datetime.fromisoformat(value.strip().replace(' ', 'T'))
        except ValueError:
            return None
    return None


def _kind(values):
    """number, date, boolean or text: a column takes a type only when every filled cell has it."""
    filled = [v for v in values if not _empty(v)]
    if not filled:
        return 'text'
    if all(isinstance(v, bool) for v in filled):
        return 'boolean'
    if all(_number(v) is not None for v in filled):
        return 'number'
    if all(_date(v) is not None for v in filled):
        return 'date'
    return 'text'


def _display(value):
    if _empty(value):
        return ''
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if (value.hour, value.minute, value.second) == (0, 0, 0) else value.isoformat(sep=' ', timespec='minutes')
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    text = ' '.join(str(value).split())
    return text if len(text) <= MAX_CELL else text[:MAX_CELL - 1] + '…'


def _sheet(name, rows, notes):
    rows, width = _trim(rows)
    if not rows or not width:
        return None
    if width > MAX_COLUMNS:
        raise SpreadsheetError(f'Sheet "{name}" has {width} columns; Lapi reads up to {MAX_COLUMNS}.')
    start = next(i for i, row in enumerate(rows) if not all(_empty(v) for v in row))
    if _looks_like_header(rows[start], rows[start + 1:]):
        columns, data, first_row = _headers(rows[start], width), rows[start + 1:], start + 2
    else:
        columns, data, first_row = [f'Column {_column_letter(i)}' for i in range(width)], rows[start:], start + 1
        notes = notes + ['no header row found: columns are named by letter']
    numbered = [(first_row + i, row) for i, row in enumerate(data) if not all(_empty(v) for v in row)]
    kinds = [_kind([row[i] for _, row in numbered]) for i in range(width)]
    return {'name': str(name)[:100], 'columns': columns, 'kinds': kinds, 'rows': numbered, 'notes': notes}


def read(name, raw):
    """Every non-empty sheet as {'name', 'columns', 'kinds', 'rows': [(row number, values)], 'notes'}."""
    if name.lower().endswith('.xlsx'):
        found = _xlsx_sheets(raw)
    else:
        found = [(name.rsplit('.', 1)[0][:100] or 'Sheet1', _csv_rows(name, raw), [])]
    sheets = [s for s in (_sheet(n, rows, notes) for n, rows, notes in found) if s]
    if not sheets:
        raise SpreadsheetError('The spreadsheet has no data.')
    total = sum(len(s['rows']) for s in sheets)
    if total > MAX_ROWS:
        raise SpreadsheetError(f'The spreadsheet has {total:,} data rows; Lapi reads up to {MAX_ROWS:,}. Split it into smaller files.')
    return sheets


# --- Search text -----------------------------------------------------------------------------------

def record(sheet, number, values):
    """One row as a self-contained line: every value keeps its column name and the row number."""
    cells = [f'{column}: {_display(value)}' for column, value in zip(sheet['columns'], values) if not _empty(value)]
    # Comments, links, colours and items attached to this row (workbook.attach_notes) stay with it.
    cells += sheet.get('row_notes', {}).get(number, [])
    return f'Row {number} · ' + ' · '.join(cells) if cells else ''


def header(sheet):
    return f'Sheet "{sheet["name"]}" · columns: ' + ', '.join(sheet['columns'])


def summary(sheet):
    """An overview of the sheet: size, and per column its type, filled count and range or examples."""
    lines = [f'Sheet "{sheet["name"]}" summary: {len(sheet["rows"])} data rows, {len(sheet["columns"])} columns.']
    for i, (column, kind) in enumerate(zip(sheet['columns'], sheet['kinds'])):
        values = [row[i] for _, row in sheet['rows'] if not _empty(row[i])]
        if kind == 'number':
            numbers = [_number(v) for v in values]
            detail = f'number, {len(numbers)} filled, min {_display(min(numbers))}, max {_display(max(numbers))}, total {_display(round(sum(numbers), 6))}' if numbers else 'number, empty'
        elif kind == 'date':
            dates = sorted(_date(v) for v in values)
            detail = f'date, {len(dates)} filled, from {_display(dates[0])} to {_display(dates[-1])}' if dates else 'date, empty'
        else:
            distinct = list(dict.fromkeys(_display(v) for v in values))
            detail = f'{kind}, {len(values)} filled, {len(distinct)} different' + (', e.g. ' + '; '.join(distinct[:SUMMARY_EXAMPLES]) if distinct else '')
        lines.append(f'- {column}: {detail}')
    if sheet['notes']:
        lines.append('Notes: ' + '; '.join(sheet['notes']) + '.')
    return '\n'.join(lines)


def pages(sheets):
    """Pages for the indexing pipeline: per sheet, a summary page and a table page whose rows are
    packed by `pack` (not cut by characters)."""
    out = []
    for index, sheet in enumerate(sheets, start=1):
        rows = [(n, record(sheet, n, values)) for n, values in sheet['rows']]
        rows = [(n, text) for n, text in rows if text]
        base = {'page': index, 'sheet': sheet['name']}
        out.append(base | {'text': summary(sheet), 'section': f'Sheet: {sheet["name"]} · summary', 'table_summary': True})
        if rows:
            out.append(base | {'text': header(sheet) + '\n' + '\n'.join(text for _, text in rows),
                               'section': f'Sheet: {sheet["name"]}', 'table_header': header(sheet), 'table_rows': rows})
    return out


def _split_row(head, number, text, fits):
    """A row too long for one chunk: split between "Column: value" pairs, never inside one."""
    prefix = f'Row {number} · '
    cells = text[len(prefix):].split(' · ') if text.startswith(prefix) else [text]
    parts, current = [], []
    for cell in cells:
        candidate = current + [cell]
        if current and not fits(head + '\n' + f'Row {number} (part {len(parts) + 1}) · ' + ' · '.join(candidate)):
            parts.append(current)
            current = [cell]
        else:
            current = candidate
    if current:
        parts.append(current)
    out = []
    for i, cells in enumerate(parts, start=1):
        line = f'Row {number} (part {i}) · ' + ' · '.join(cells)
        while not fits(head + '\n' + line) and len(line) > 40:
            line = line[:int(len(line) * 0.8)].rstrip() + '…'  # one huge cell: keep what fits
        out.append(line)
    return out


def pack(head, rows, fits):
    """Chunks of whole rows: [(text, first row, last row)], each starting with the sheet header."""
    chunks, current = [], []
    def flush():
        if current:
            chunks.append((head + '\n' + '\n'.join(text for _, text in current), current[0][0], current[-1][0]))
    for number, text in rows:
        if fits(head + '\n' + '\n'.join([t for _, t in current] + [text])):
            current.append((number, text))
            continue
        flush()
        current = []
        if fits(head + '\n' + text):
            current = [(number, text)]
        else:
            for part in _split_row(head, number, text, fits):
                chunks.append((head + '\n' + part, number, number))
    flush()
    return chunks


# --- Exact table: storage and safe calculations -----------------------------------------------------

def _arrow(sheet):
    import pyarrow as pa
    columns = {}
    for i, (column, kind) in enumerate(zip(sheet['columns'], sheet['kinds'])):
        values = [row[i] for _, row in sheet['rows']]
        if kind == 'number':
            columns[column] = pa.array([None if _empty(v) else _number(v) for v in values], pa.float64())
        elif kind == 'date':
            columns[column] = pa.array([None if _empty(v) else _date(v) for v in values], pa.timestamp('s'))
        elif kind == 'boolean':
            columns[column] = pa.array([None if _empty(v) else bool(v) for v in values], pa.bool_())
        else:
            columns[column] = pa.array([None if _empty(v) else _display(v) for v in values], pa.string())
    columns['__row'] = pa.array([n for n, _ in sheet['rows']], pa.int64())
    return pa.table(columns)


def manifest_key(document_id):
    return f'documents/{document_id}/tables.json'


def save(store, document_id, sheets):
    """Store each sheet as Parquet plus a manifest of sheets, columns and types."""
    import pyarrow.parquet as pq
    entries = []
    for index, sheet in enumerate(sheets, start=1):
        buffer = io.BytesIO()
        pq.write_table(_arrow(sheet), buffer)
        key = f'documents/{document_id}/tables/{index}.parquet'
        store.put(key, buffer.getvalue())
        entries.append({'index': index, 'name': sheet['name'], 'key': key, 'rows': len(sheet['rows']),
                        'columns': [{'name': c, 'type': k} for c, k in zip(sheet['columns'], sheet['kinds'])], 'notes': sheet['notes']})
    store.put_json(manifest_key(document_id), {'sheets': entries})
    return entries


def tables(store, document_id):
    try:
        return store.json(manifest_key(document_id))['sheets']
    except (FileNotFoundError, KeyError, OSError):
        return []


OPS = ('=', '!=', '>', '>=', '<', '<=', 'contains', 'in', 'between', 'is empty', 'is not empty')
METRICS = ('count', 'sum', 'avg', 'min', 'max', 'count_distinct')
BUCKETS = ('day', 'month', 'quarter', 'year')
MAX_RESULT_ROWS = 200


def _column(entry, name):
    for column in entry['columns']:
        if column['name'].lower() == str(name).strip().lower():
            return column
    raise SpreadsheetError(f'Sheet "{entry["name"]}" has no column "{name}". Columns: ' + ', '.join(c['name'] for c in entry['columns']))


def _scalar(column, value):
    if column['type'] == 'number':
        number = _number(value)
        if number is None:
            raise SpreadsheetError(f'"{value}" is not a number for column "{column["name"]}".')
        return number
    if column['type'] == 'date':
        date = _date(value) if not isinstance(value, str) or DATE.match(value.strip()) else None
        if date is None:
            raise SpreadsheetError(f'Use a date like 2026-10-01 for column "{column["name"]}".')
        return date
    if column['type'] == 'boolean':
        return str(value).strip().lower() in ('true', 'yes', '1')
    return str(value)


def _mask(table, entry, condition):
    import pyarrow as pa
    import pyarrow.compute as pc
    column = _column(entry, condition.get('column'))
    data = table[column['name']]
    op = str(condition.get('op', '=')).strip().lower()
    if op not in OPS:
        raise SpreadsheetError(f'Unknown filter "{op}". Use one of: ' + ', '.join(OPS))
    if op == 'is empty':
        return pc.is_null(data) if column['type'] != 'text' else pc.or_(pc.is_null(data), pc.equal(data, ''))
    if op == 'is not empty':
        return pc.is_valid(data) if column['type'] != 'text' else pc.and_(pc.is_valid(data), pc.not_equal(data, ''))
    value = condition.get('value')
    if op == 'contains':
        return pc.match_substring(pc.cast(data, pa.string()), str(value), ignore_case=True)
    if op == 'in':
        values = value if isinstance(value, list) else [value]
        if column['type'] == 'text':
            return pc.is_in(pc.utf8_lower(data), value_set=pa.array([str(v).lower() for v in values]))
        return pc.is_in(data, value_set=pa.array([_scalar(column, v) for v in values], type=data.type))
    if op == 'between':
        low, high = (value if isinstance(value, list) and len(value) == 2 else (None, None))
        if low is None or high is None:
            raise SpreadsheetError('"between" needs two values: [from, to].')
        return pc.and_(pc.greater_equal(data, pa.scalar(_scalar(column, low), data.type)), pc.less_equal(data, pa.scalar(_scalar(column, high), data.type)))
    scalar = _scalar(column, value)
    if column['type'] == 'text' and op in ('=', '!='):
        equal = pc.equal(pc.utf8_lower(data), str(scalar).lower())
        return equal if op == '=' else pc.invert(equal)
    compare = {'=': pc.equal, '!=': pc.not_equal, '>': pc.greater, '>=': pc.greater_equal, '<': pc.less, '<=': pc.less_equal}[op]
    return compare(data, pa.scalar(scalar, data.type))


def _bucket(table, column, bucket):
    import pyarrow.compute as pc
    data = table[column['name']]
    if column['type'] != 'date':
        raise SpreadsheetError(f'"{column["name"]}" is not a date column, so it cannot be grouped by {bucket}.')
    if bucket == 'day':
        return pc.strftime(data, format='%Y-%m-%d')
    if bucket == 'month':
        return pc.strftime(data, format='%Y-%m')
    if bucket == 'year':
        return pc.strftime(data, format='%Y')
    quarter = pc.cast(pc.quarter(data), 'string')
    return pc.binary_join_element_wise(pc.strftime(data, format='%Y'), pc.binary_join_element_wise('Q', quarter, ''), '-')


def query(store, document_id, spec):
    """Run one validated calculation on a stored sheet.

    spec: {"sheet": name or index (optional), "filters": [{"column", "op", "value"}],
           "group_by": [column or {"column", "bucket": day|month|quarter|year}] (max 2),
           "metrics": [{"op": count|sum|avg|min|max|count_distinct, "column"}],
           "columns": [...] (list matching rows instead of aggregating), "sort": {"by", "desc"}, "limit"}
    Returns {"sheet", "columns", "rows", "matched_rows", "row_numbers", "truncated"}.
    """
    import pyarrow as pa
    import pyarrow.compute as pc
    import pyarrow.parquet as pq
    entries = tables(store, document_id)
    if not entries:
        raise SpreadsheetError('This document is not a spreadsheet, or it has not been read yet.')
    wanted = spec.get('sheet')
    if wanted in (None, ''):
        entry = entries[0]
    else:
        entry = next((e for e in entries if str(e['index']) == str(wanted) or e['name'].lower() == str(wanted).strip().lower()), None)
        if entry is None:
            raise SpreadsheetError('No sheet named "' + str(wanted) + '". Sheets: ' + ', '.join(e['name'] for e in entries))
    table = pq.read_table(io.BytesIO(store.get(entry['key'])))
    filters = spec.get('filters') or []
    if not isinstance(filters, list) or len(filters) > 10:
        raise SpreadsheetError('Use up to 10 filters.')
    mask = None
    for condition in filters:
        part = _mask(table, entry, condition if isinstance(condition, dict) else {})
        mask = part if mask is None else pc.and_(mask, part)
    if mask is not None:
        table = table.filter(pc.fill_null(mask, False))
    matched = table.num_rows
    limit = max(1, min(int(spec.get('limit') or 50), MAX_RESULT_ROWS))
    group_by = spec.get('group_by') or []
    metrics = spec.get('metrics') or []
    if not isinstance(group_by, list) or len(group_by) > 2:
        raise SpreadsheetError('Group by at most 2 columns.')
    if not isinstance(metrics, list) or len(metrics) > 6:
        raise SpreadsheetError('Use up to 6 calculations.')
    if group_by or metrics:
        keys, names = [], []
        for item in group_by:
            column = _column(entry, item.get('column') if isinstance(item, dict) else item)
            bucket = item.get('bucket') if isinstance(item, dict) else None
            if bucket:
                if bucket not in BUCKETS:
                    raise SpreadsheetError('Group dates by day, month, quarter or year.')
                label = f'{column["name"]} ({bucket})'
                table = table.append_column(label, _bucket(table, column, bucket))
            else:
                label = column['name']
            keys.append(label); names.append(label)
        aggregations, labels = [], []
        for metric in metrics or [{'op': 'count'}]:
            op = str(metric.get('op', 'count')).lower()
            if op not in METRICS:
                raise SpreadsheetError('Calculations: ' + ', '.join(METRICS))
            if op == 'count' and not metric.get('column'):
                aggregations.append(('__row', 'count')); labels.append('Count')
                continue
            column = _column(entry, metric.get('column'))
            if op in ('sum', 'avg') and column['type'] != 'number':
                raise SpreadsheetError(f'"{column["name"]}" is not a number column, so it cannot be summed or averaged.')
            arrow_op = {'avg': 'mean', 'count_distinct': 'count_distinct'}.get(op, op)
            aggregations.append((column['name'], arrow_op)); labels.append(f'{op} of {column["name"]}')
        if keys:
            result = table.group_by(keys).aggregate(aggregations)
            produced = [f'{c}_{o}' for c, o in aggregations]
            result = result.select(keys + produced).rename_columns(names + labels)
        else:
            values = []
            for (column, op) in aggregations:
                data = table[column]
                values.append(pc.count(data).as_py() if op == 'count' else pc.count_distinct(data).as_py() if op == 'count_distinct' else getattr(pc, op)(data).as_py())
            result = pa.table({label: [value] for label, value in zip(labels, values)})
        row_numbers = []
    else:
        chosen = [_column(entry, c)['name'] for c in (spec.get('columns') or [])] or [c['name'] for c in entry['columns']]
        row_numbers = table['__row'].to_pylist()[:limit]
        result = table.select(chosen)
    sort = spec.get('sort') or {}
    if isinstance(sort, dict) and sort.get('by'):
        by = next((n for n in result.column_names if n.lower() == str(sort['by']).strip().lower()), None)
        if by is None:
            raise SpreadsheetError(f'Cannot sort by "{sort["by"]}". Sort by one of: ' + ', '.join(result.column_names))
        result = result.sort_by([(by, 'descending' if sort.get('desc') else 'ascending')])
    truncated = result.num_rows > limit
    result = result.slice(0, limit)
    rows = [[round(v, 6) if isinstance(v, float) else v if isinstance(v, int) and not isinstance(v, bool) else _display(v) for v in row.values()]
            for row in result.to_pylist()]
    return {'document_id': document_id, 'sheet': entry['name'], 'columns': result.column_names, 'rows': rows,
            'matched_rows': matched, 'row_numbers': row_numbers, 'truncated': truncated}
